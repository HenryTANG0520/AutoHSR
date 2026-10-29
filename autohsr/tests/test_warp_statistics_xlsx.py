"""Synthetic window-statistics and workbook contracts, without real records."""
from contextlib import redirect_stdout,redirect_stderr
from copy import deepcopy
import hashlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES
from openpyxl.worksheet._write_only import WriteOnlyWorksheet

from autohsr.warp_records import WarpError,export
from autohsr.warp_fetch import WarpStopped
from autohsr.warp_statistics import analyze,positions,public_statistics
from autohsr.warp_xlsx import write_new
from autohsr.warp_service import run


def row(identity,rank='3',pool='11',count='1',name='虚构角色'):
    value=dict(id=str(identity),gacha_id='1001',gacha_type=pool,item_id='1000',
               time='2026-10-29 12:34:56',name=name,item_type='角色')
    if rank is not None:value['rank_type']=rank
    if count is not None:value['count']=count
    return value


def group(rows=(),uid='100000000'):
    return dict(uid=uid,timezone=8,lang='zh-cn',list=list(rows))


def sample():return group([row(i+1,rank,count='3' if i==0 else '1')
                           for i,rank in enumerate(('3','5','4','3','5',None,'5','3','5','3'))])


class StatisticsTests(unittest.TestCase):
    def test_left_prefix_unknown_rank_and_tail_are_not_actual_pity(self):
        data=sample();before=deepcopy(data);stats=analyze([data]);pool=stats['accounts'][0]['pools']['11']
        self.assertEqual(pool['rank_counts'],{'3':4,'4':1,'5':4,'unknown':1})
        self.assertEqual(pool['item_quantity_sum'],12);self.assertEqual(pool['records'],10)
        self.assertEqual((pool['bounded_intervals'],pool['mean_bounded_interval_records']),(2,2.5))
        self.assertEqual((pool['observed_tail_records'],pool['known_tail_records']),(1,1))
        self.assertEqual([r['bounded_interval_records'] for r in pool['five_star_history']],[None,3,None,2])
        self.assertEqual(pool['rank_shares']['5'],.4);self.assertFalse(stats['history_complete'])
        self.assertEqual(data,before)

    def test_empty_no_five_missing_count_and_unknown_tail_remain_explicit(self):
        for rows in ([],[row(1)],[row(1,'5'),row(2,None,count=None)]):
            pool=analyze([group(rows)])['accounts'][0]['pools']['11']
            self.assertIsNone(pool['mean_bounded_interval_records'])
            self.assertIsNone(pool['known_tail_records'])
            if not rows:self.assertIsNone(pool['rank_shares']['5']);self.assertIsNone(pool['range_start'])
        pool=analyze([group([row(1,'5'),row(2,None,count=None)])])['accounts'][0]['pools']['11']
        self.assertIsNone(pool['item_quantity_sum']);self.assertEqual(pool['observed_tail_records'],1)

    def test_accounts_pools_numeric_ids_and_time_disagreement_stay_separate(self):
        a=group([row(10,'5'),row(9,'5'),row(11,'3','1'),row(12,'5')])
        a['list'][0]['time']='2026-10-28 00:00:00'
        b=group([row(1,'3'),row(2,'5')],uid='200000000')
        stats=analyze([a,b]);pool=stats['accounts'][0]['pools']['11']
        self.assertEqual([r['id'] for r in pool['five_star_history']],['9','10','12'])
        self.assertEqual(pool['mean_bounded_interval_records'],1)
        self.assertFalse(pool['time_order_consistent'])
        self.assertEqual(stats['accounts'][0]['pools']['1']['records'],1)
        self.assertIsNone(stats['accounts'][1]['pools']['11']['mean_bounded_interval_records'])

    def test_public_statistics_is_a_detached_allowlist_without_identifiers(self):
        private=analyze([sample()]);public=public_statistics(private);text=json.dumps(public,ensure_ascii=False)
        for secret in ('100000000','虚构角色','2026-10-29','five_star_history','"uid"'):
            self.assertNotIn(secret,text)
        public['accounts'][0]['pools']['11']['rank_counts']['5']=999
        self.assertEqual(private['accounts'][0]['pools']['11']['rank_counts']['5'],4)

    def test_positions_do_not_expand_item_quantity_into_record_count(self):
        rows=[row(1,'5',count='3'),row(2,'3',count='2'),row(3,'5',count=None)]
        values=list(positions(rows))
        self.assertEqual([v[1]['since_previous_five_records'] for v in values],[None,1,2])
        self.assertEqual([v[1]['index'] for v in values],[1,2,3])

    def test_cancellation_before_analysis_does_not_collect_or_invent_a_result(self):
        with self.assertRaises(WarpStopped):analyze([sample()],cancel=lambda:True)


class WorkbookTests(unittest.TestCase):
    def test_workbook_styles_ids_unknowns_and_filters_are_valid(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'records.xlsx';a=sample();a['list'][0]['id']='9007199254740993001'
            for i,value in enumerate(a['list']):value['id']=str(9007199254740993001+i)
            a['list'][0]['name']='=HYPERLINK("https://example.invalid","literal")'
            a['uid']='10000000000000000001';original=deepcopy(a)
            checksum=write_new(path,[a]);self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),checksum)
            wb=load_workbook(path,data_only=False);self.addCleanup(wb.close)
            self.assertEqual(wb.sheetnames,['分池统计','记录-10000000000000000001'])
            ws=wb.worksheets[1]
            self.assertEqual(ws['B2'].value,a['uid']);self.assertEqual(ws['B2'].data_type,'s')
            self.assertEqual(ws['K6'].value,'9007199254740993001');self.assertEqual(ws['K6'].data_type,'s')
            self.assertEqual(ws['B6'].data_type,'s');self.assertTrue(ws['B6'].value.startswith('=HYPERLINK'))
            self.assertEqual(ws['F6'].value,3);self.assertEqual(ws['G6'].value,1)
            self.assertIsNone(ws['I7'].value);self.assertEqual(ws['I10'].value,3)
            self.assertIsNone(ws['D11'].value);self.assertIsNone(ws['I12'].value)
            self.assertEqual(ws.freeze_panes,'A6');self.assertEqual(ws.auto_filter.ref,'A5:M15')
            self.assertTrue(ws['A7'].font.color.rgb.endswith('9A6700'))
            self.assertTrue(ws['A8'].font.color.rgb.endswith('6F42C1'))
            self.assertIn('B3:M3',str(ws.merged_cells));self.assertEqual(ws['A6'].value,'2026-10-29 12:34:56')
            stats=wb.worksheets[0];self.assertEqual(stats['C8'].value,10);self.assertEqual(stats['K8'].value,2.5)
            self.assertEqual(stats['H8'].number_format,'0.00%');self.assertEqual(stats.freeze_panes,'C6')
            self.assertEqual(a,original)
            with zipfile.ZipFile(path) as archive:
                for name in archive.namelist():
                    if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                        xml=ET.fromstring(archive.read(name))
                        self.assertEqual(xml.findall('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}f'),[])
                self.assertFalse(any('externalLink' in name or 'vbaProject' in name for name in archive.namelist()))

    def test_multiple_and_empty_accounts_have_distinct_sheets_and_no_first_uid_fallback(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source=root/'input.json';source.write_text(json.dumps(export([sample(),group([],uid='200000000')])),encoding='utf8')
            output=root/'selected.xlsx';result=run(imports=[source],output=output,format='xlsx',selected_uid='200000000')
            self.assertEqual((result['accounts'],result['records'],result['format']),(1,0,'xlsx'))
            wb=load_workbook(output);self.addCleanup(wb.close);self.assertEqual(wb.sheetnames,['分池统计','记录-200000000'])
            self.assertEqual(wb.worksheets[1]['B2'].value,'200000000');self.assertEqual(wb.worksheets[1].max_row,5)
            other=root/'both.xlsx';write_new(other,[sample(),group([],uid='200000000')])
            wb2=load_workbook(other);self.addCleanup(wb2.close);self.assertEqual(len(wb2.sheetnames),3)

    def test_analysis_only_cli_does_not_write_or_import_excel_or_game_settings(self):
        from autohsr.__main__ import main
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source=root/'input.json';source.write_text(json.dumps(export([sample()])),encoding='utf8')
            before=source.read_bytes();output=io.StringIO();events=[]
            with patch('autohsr.config.Settings.load',side_effect=AssertionError('no game configuration')), \
                 patch('autohsr.warp_xlsx._write',side_effect=AssertionError('no workbook')),redirect_stdout(output):
                self.assertEqual(main(['--warp-analysis','--warp-import',str(source)]),0)
            public=json.loads(output.getvalue());self.assertEqual(public['mode'],'analysis')
            self.assertEqual(public['statistics']['accounts'][0]['pools']['11']['mean_bounded_interval_records'],2.5)
            self.assertNotIn('100000000',output.getvalue());self.assertNotIn('虚构角色',output.getvalue())
            result=run(imports=[source],analysis_only=True,emit=events.append)
            self.assertEqual(result['mode'],'analysis');self.assertEqual(events[-1]['type'],'warp_analysis_finished')
            self.assertNotIn('output_sha256',result);self.assertEqual(list(root.iterdir()),[source]);self.assertEqual(source.read_bytes(),before)

    def test_cancel_mid_stream_removes_owned_sheet_xml_and_unpublished_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);path=root/'cancel.xlsx';original=WriteOnlyWorksheet.append;stopped=[False]
            baseline=set(ALL_TEMP_FILES)
            def append(ws,values):
                result=original(ws,values)
                if ws.title.startswith('记录-'):stopped[0]=True
                return result
            with patch.object(WriteOnlyWorksheet,'append',append),self.assertRaises(WarpStopped):
                write_new(path,[group([row(i+1) for i in range(140)])],cancel=lambda:stopped[0])
            self.assertFalse(path.exists());self.assertEqual(list(root.iterdir()),[]);self.assertEqual(set(ALL_TEMP_FILES),baseline)

    def test_failure_and_invalid_xml_preserve_input_and_cleanup(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);path=root/'bad.xlsx';baseline=set(ALL_TEMP_FILES)
            a=group([row(1,name='bad\x00text')]);before=deepcopy(a)
            with self.assertRaises(WarpError):write_new(path,[a])
            self.assertFalse(path.exists());self.assertEqual(a,before);self.assertEqual(set(ALL_TEMP_FILES),baseline)
            with patch('autohsr.warp_service.os.fsync',side_effect=OSError('synthetic disk failure')),self.assertRaises(OSError):
                write_new(path,[sample()])
            self.assertFalse(path.exists());self.assertEqual(list(root.iterdir()),[])

    def test_existing_output_and_wrong_extension_are_rejected_before_publication(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);path=root/'out.xlsx';path.write_bytes(b'old bytes')
            with self.assertRaises(WarpError):write_new(path,[sample()])
            self.assertEqual(path.read_bytes(),b'old bytes')
            with self.assertRaises(WarpError):write_new(root/'bad.json',[sample()])

    def test_cli_analysis_export_scope_and_xlsx_source_are_explicit(self):
        from autohsr.__main__ import main
        for args in (['--warp-analysis','--warp-output','unused.xlsx'],['--warp-analysis','--warp-format','xlsx'],
                     ['--warp-analysis','--warp-records'],['--warp-analysis','--config','unused.json']):
            with redirect_stderr(io.StringIO()),self.assertRaises(SystemExit):main(args)
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source=root/'source.json';source.write_text(json.dumps(export([sample()])),encoding='utf8')
            out=root/'out.xlsx';stdout=io.StringIO()
            with redirect_stdout(stdout):
                self.assertEqual(main(['--warp-records','--warp-import',str(source),'--warp-format','xlsx','--warp-output',str(out)]),0)
            self.assertEqual(json.loads(stdout.getvalue())['format'],'xlsx');self.assertTrue(zipfile.is_zipfile(out))
            self.assertNotIn('100000000',stdout.getvalue())

    def test_missing_excel_engine_fails_before_source_reads_and_json_still_works(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source=root/'input.json';source.write_text(json.dumps(export([sample()])),encoding='utf8')
            with patch('autohsr.warp_xlsx.require_engine',side_effect=WarpError('dependency unavailable')):
                with self.assertRaises(WarpError):run(url_file=root/'nonexistent-private-link.txt',output=root/'out.xlsx',format='xlsx')
                result=run(imports=[source],output=root/'out.json')
                self.assertEqual(result['format'],'uigf-v4.2')
                self.assertEqual(run(imports=[source],analysis_only=True)['mode'],'analysis')
            self.assertFalse((root/'out.xlsx').exists())
            with patch('openpyxl.__version__','0.0'),self.assertRaises(WarpError):
                run(imports=[source],output=root/'version.xlsx',format='xlsx')
            self.assertFalse((root/'version.xlsx').exists())


if __name__=='__main__':unittest.main()
