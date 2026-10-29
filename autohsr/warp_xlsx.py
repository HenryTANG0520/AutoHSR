"""Styled, streamed record workbooks; IDs/text never become Excel formulas.

Behaviour reference: March7thAssistant __onExportExcelBtnClicked (7cf6c6b).
The workbook is a reading view; original JSON remains the exchange archive.
"""
from pathlib import Path
import re

from .warp_records import POOLS,WarpError,merge
from .warp_fetch import checkpoint
from .warp_statistics import POOL_NAMES,analyze,positions

SCOPE='仅当前记录窗口；不代表完整历史或实时保底'
INVALID_XML=re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')
RECORD_HEADERS=('时间','名称','类别','星级','卡池','物品数量','账号窗口序号','本池窗口序号',
                '已知五星后记录数','计数范围','记录ID','物品ID','卡池ID')
STATS_HEADERS=('账号','卡池','记录数','3星记录','4星记录','5星记录','未知星级','5星占比',
               '已观察尾段记录','星级已知尾段','五星间平均记录数','观察区间样本数','时间顺序','窗口起始','窗口结束')
SCOPE_LABELS={'observed':'仅已观察记录','window_prefix':'窗口前缀，未见上个五星','unknown_rank':'缺失星级，间隔未知'}


def require_engine():
    try:import openpyxl
    except ImportError as error:raise WarpError('Excel导出缺少openpyxl，请使用项目setup安装当前依赖。') from error
    if openpyxl.__version__!='3.1.5':raise WarpError('Excel导出依赖版本不符，请使用项目setup安装固定依赖。')


def _write(stream,values,*,cancel=None):
    require_engine()
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.utils import get_column_letter
    statistics=analyze(values,history=False,cancel=cancel)
    workbook=Workbook(write_only=True);workbook.properties.creator='AutoHSR'
    workbook.properties.title='跃迁记录：已导入窗口'
    fonts={key:Font(name='Microsoft YaHei',size=10,color=color,bold=key in ('header','5','title'))
           for key,color in {'header':'FFFFFF','5':'9A6700','4':'6F42C1','body':'263238','title':'263238'}.items()}
    fills={key:PatternFill('solid',fgColor=color) for key,color in
           {'header':'263238','5':'FFF4D6','4':'F4EDFA','even':'F4F6F8','body':'FFFFFF','title':'E7ECF0'}.items()}

    def cell(sheet,value,style='body',number_format=None):
        if isinstance(value,str) and INVALID_XML.search(value):raise WarpError('记录文字含Excel无法表示的控制字符，请保留JSON原件。')
        result=WriteOnlyCell(sheet,value=value)
        if isinstance(value,str):result.data_type='s';result.number_format='@'
        elif number_format is not None:result.number_format=number_format
        result.font=fonts.get(style,fonts['body']);result.fill=fills.get(style,fills['body'])
        result.alignment=Alignment(vertical='center',wrap_text=True)
        return result

    def sheet(name,headers,widths,metadata,*,freeze='A6'):
        ws=workbook.create_sheet(name);last=get_column_letter(len(headers))
        ws.freeze_panes=freeze;ws.sheet_view.showGridLines=False
        for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
        ws.merged_cells.add(f'A1:{last}1');ws.merged_cells.add(f'B3:{last}3')
        for number,height in ((1,26),(2,24),(3,32),(5,44)):ws.row_dimensions[number].height=height
        ws.print_title_rows='1:5';ws.page_setup.orientation='landscape'
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.append([cell(ws,'AutoHSR 跃迁记录窗口','title')])
        ws.append([cell(ws,v) for v in metadata])
        ws.append([cell(ws,'统计范围'),cell(ws,SCOPE)])
        ws.append([]);ws.append([cell(ws,v,'header') for v in headers])
        return ws

    try:
        stats=sheet('分池统计',STATS_HEADERS,(24,22,11,11,11,11,12,12,18,18,20,18,12,23,23),
                    ('账号数',len(values),None,'记录数',sum(len(a['list']) for a in values)),freeze='C6')
        for account in statistics['accounts']:
            checkpoint(cancel)
            for key in POOLS:
                data=account['pools'][key];counts=data['rank_counts']
                row=(account['uid'],POOL_NAMES[key],data['records'],counts['3'],counts['4'],counts['5'],counts['unknown'],
                     data['rank_shares']['5'],data['observed_tail_records'],data['known_tail_records'],
                     data['mean_bounded_interval_records'],data['bounded_intervals'],
                     ('未见逆序' if data['time_order_consistent'] else '发现逆序') if data['records'] else None,
                     data['range_start'],data['range_end'])
                cells=[cell(stats,v,'even' if (account['index']+POOLS.index(key))%2 else 'body') for v in row]
                cells[7].number_format='0.00%';cells[10].number_format='0.00';stats.append(cells)
        stats.auto_filter.ref=f'A5:O{5+6*len(values)}'
        for value in values:
            checkpoint(cancel)
            ws=sheet('记录-'+value['uid'],RECORD_HEADERS,(23,28,14,9,23,12,15,15,22,28,25,20,20),
                     ('账号',value['uid'],None,'语言',value.get('lang','未知'),None,'时区',f'UTC{value["timezone"]:+03d}:00'))
            for row,position in positions(value['list']):
                if position['index']%128==1:checkpoint(cancel)
                rank=row.get('rank_type')
                style=rank if rank in ('4','5') else 'even' if position['index']%2==0 else 'body'
                data=(row['time'],row.get('name'),row.get('item_type'),int(rank) if rank else None,
                      POOL_NAMES[row['gacha_type']],int(row['count']) if 'count' in row else None,
                      position['index'],position['pool_index'],position['since_previous_five_records'],
                      SCOPE_LABELS[position['scope']],row['id'],row['item_id'],row['gacha_id'])
                ws.append([cell(ws,v,style) for v in data])
            ws.auto_filter.ref=f'A5:M{5+len(value["list"])}'
        checkpoint(cancel);workbook.save(stream)
    finally:
        # Streaming sheet XML belongs to this workbook. Pinned openpyxl 3.1.5
        # exposes no public abort cleanup, so release its owned temporary files.
        for ws in workbook.worksheets:
            writer=getattr(ws,'_writer',None)
            if writer is not None:
                try:
                    if not ws.closed:ws.close()
                finally:
                    if Path(writer.out).is_file():writer.cleanup()
        workbook.close()


def write_new(path,accounts,*,cancel=None):
    from .warp_service import publish_new
    path=Path(path)
    if path.suffix.lower()!='.xlsx':raise WarpError('Excel导出文件应使用.xlsx扩展名。')
    checkpoint(cancel);values=merge(accounts)
    if not values:raise WarpError('没有可导出的账号记录。')
    return publish_new(path,lambda stream:_write(stream,values,cancel=cancel),cancel=cancel,limit=128*1024*1024)
